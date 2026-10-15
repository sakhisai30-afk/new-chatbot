"""Creates ZuiGO-translation-check.xlsx for Google Sheets cross-checking.
Run: python3 scripts/i18n/make-workbook.py   (needs openpyxl)
Upload the .xlsx to Google Sheets (File > Import) so GOOGLETRANSLATE works."""
import json, pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

here = pathlib.Path(__file__).parent
en = json.loads((here / 'en.flat.json').read_text(encoding='utf-8'))
LANGS = [('de', 'German', 'de'), ('zh', 'Chinese', 'zh-CN'), ('ko', 'Korean', 'ko'), ('ja', 'Japanese', 'ja'),
         ('ru', 'Russian', 'ru'), ('es', 'Spanish', 'es'), ('fr', 'French', 'fr')]
head = PatternFill('solid', fgColor='1F2A44')
wb = Workbook()
ws = wb.active; ws.title = 'READ ME'
for i, line in enumerate([
    'ZuiGO chatbot - translation cross-check',
    '',
    '1. Upload this file to Google Sheets (File > Import > Upload). Do NOT use Excel: GOOGLETRANSLATE only exists in Google Sheets.',
    '2. Each language tab has: key | English | Translation used in the app | Google Translate of English | Google Translate of the app text back to English.',
    '3. Compare column C (used in app) with D (Google). Wording may differ - that is normal. Meaning, names and numbers must match.',
    '4. Column E (back to English) should read like the English in column B. If it does not, the translation is probably wrong.',
    '5. Mark column F: OK / Fix / Native-check. Put a corrected text in column G for any Fix, then copy it into scripts/i18n/translations/<code>.json',
    '6. Then run: npm run i18n:build && npm run i18n:validate',
    '',
    'Protected (must stay unchanged): ZuiGO, product/tech names (React, AWS, Docker...), hello@zuigo.ai, the address, {{number}} and {{serviceName}}.',
    'Note: GOOGLETRANSLATE may mangle {{number}} / {{serviceName}} in column D - ignore that, check column C.',
    'Note: the translations in column C were written by Claude (AI) and are NOT reviewed by a native speaker.'], 1):
    ws.cell(row=i, column=1, value=line)
ws['A1'].font = Font(bold=True, size=14); ws.column_dimensions['A'].width = 140

for code, name, gcode in LANGS:
    tr = json.loads((here / 'translations' / f'{code}.json').read_text(encoding='utf-8'))
    s = wb.create_sheet(f'{code} {name}')
    for c, h in enumerate(['Key', 'English (source)', f'{name} - used in app', f'Google: EN to {gcode}', 'Google: back to EN (check)',
                           'Reviewer: OK / Fix / Native-check', 'Corrected text (if Fix)'], 1):
        cell = s.cell(row=1, column=c, value=h); cell.font = Font(bold=True, color='FFFFFF'); cell.fill = head
    for r, (key, src) in enumerate(en.items(), 2):
        s.cell(row=r, column=1, value=key)
        s.cell(row=r, column=2, value=src)
        s.cell(row=r, column=3, value=tr[key])
        s.cell(row=r, column=4, value=f'=GOOGLETRANSLATE(B{r},"en","{gcode}")')
        s.cell(row=r, column=5, value=f'=GOOGLETRANSLATE(C{r},"{gcode}","en")')
        for c in range(1, 8): s.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical='top')
    for col, w in zip('ABCDEFG', [30, 50, 50, 50, 50, 16, 40]): s.column_dimensions[col].width = w
    s.freeze_panes = 'C2'
    dv = DataValidation(type='list', formula1='"OK,Fix,Native-check"', allow_blank=True)
    s.add_data_validation(dv); dv.add(f'F2:F{len(en)+1}')
wb.save(here.parent.parent / 'ZuiGO-translation-check.xlsx')
print('saved ZuiGO-translation-check.xlsx')
